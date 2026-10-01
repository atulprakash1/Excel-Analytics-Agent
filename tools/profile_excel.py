"""
profile_excel.py - measure the structure and quality of Excel/CSV inputs.

This script does the *measuring*; the Profiler agent does the *interpreting*.
It never modifies source files.

For every workbook it writes:
  profiles/<file_stem>.profile.json   machine-readable profile (the "data contract")
  profiles/<file_stem>.profile.html   human-readable profile report
and, across all profiled files:
  profiles/_relationships.json        candidate join keys between sheets/files

Caching and drift:
  If a profile already exists and the file hash is unchanged, the file is skipped.
  If the file changed, the new structure is compared with the old one; the result is
  stored under "drift", and the agent's "annotations" are carried forward so the
  business context is not lost. Use --force to re-profile everything.

Usage:
  python tools/profile_excel.py                      # all files in data/raw
  python tools/profile_excel.py data/raw/<file>.xlsx # one file
  python tools/profile_excel.py --force
  python tools/profile_excel.py --html-only          # re-render HTML after editing annotations
  python tools/profile_excel.py --quiet              # one status line per file (after approving knowledge)
  python tools/profile_excel.py data/raw/<file> --prefill
      # also draft the mechanical annotations of an unconfirmed profile: roles (suggested_role),
      # additivity guessed from column names, cleaning rules, a draft grain, single-value columns,
      # and open_questions for every guess, outlier and "is this file recurring?". Fills gaps only;
      # status stays draft - the agent adds business meaning and a person confirms.

Recurring files (logical sources in knowledge/source_registry.json):
  A NEW file that matches a registered source (e.g. bank_balances_2026_10.xlsx for pattern
  bank_balances_*.xlsx) inherits the confirmed annotations of the previous file of that source
  when its structure matches. New month columns in wide sheets are expected growth, not drift.
  If the structure differs, the new profile is marked needs_review.

Domain packs (knowledge/domains/*.md):
  Every profile gets `suggested_domains`, a ranking of packs by how well their signals match
  the file. Confirmed packs are recorded in annotations.domains (python tools/domains.py set).
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from collections import Counter
from datetime import date, datetime, time
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from tools.report_kit import Report, fmt_num, fmt_pct  # noqa: E402
import copy  # noqa: E402

from tools.knowledge import apply_catalog, match_source, source_files  # noqa: E402
from tools.domains import recorded_domains, suggest as suggest_domains  # noqa: E402

SCHEMA_VERSION = "1.0"
RAW_DIR = ROOT / "data" / "raw"
PROFILE_DIR = ROOT / "profiles"
SUPPORTED = {".xlsx", ".xlsm", ".csv"}
HEADER_SCAN_ROWS = 30
TOTAL_PATTERN = re.compile(r"(?i)\b(sub[\s-]?total|grand\s+total|total)\b")
ID_NAME_PATTERN = re.compile(r"(?i)(\bid\b|_id\b|\bcode\b|\bno\b|\bno\.|number|\bref\b|\bkey\b|\bsku\b)")
MONTHS = "jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec"
TENOR_PATTERN = re.compile(  # maturity / repricing buckets used as columns: O/N, 0-1M, 1-3M, >5Y, 5Y+
    r"(?i)^\s*(o/?n|overnight|t/?n|s/?n|[<>]=?\s*\d+\s*[dwmy]|\d+\s*[dwmy]?\s*(-|to)\s*\d+\s*[dwmy]|\d+\s*[dwmy]\+?)\s*$")


def is_period_column(name) -> bool:
    """True for time-period headers (Jan-26, Q1 2026, 2026) and tenor buckets (O/N, 1-3M, >5Y)."""
    return bool(PERIOD_PATTERN.match(str(name)) or TENOR_PATTERN.match(str(name)))


PERIOD_PATTERN = re.compile(
    rf"(?i)^\s*(({MONTHS})[a-z]*[\s\-/']*(\d{{2}}|\d{{4}})?|q[1-4][\s\-/]*(\d{{2}}|\d{{4}})?|"
    rf"(fy)?\s?(19|20)\d{{2}}|\d{{4}}[\-/]\d{{1,2}}|\d{{1,2}}[\-/]\d{{4}})\s*$")


# --------------------------------------------------------------------------- helpers

def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def jsonable(v):
    if v is None:
        return None
    if isinstance(v, (datetime, date, time, pd.Timestamp)):
        return v.isoformat()
    if hasattr(v, "item"):  # numpy scalar
        v = v.item()
    if isinstance(v, float) and v != v:
        return None
    return v


def to_number(s):
    if isinstance(s, str):
        t = s.strip().replace(",", "").replace("£", "").replace("$", "").replace("€", "")
        if t.endswith("%"):
            t = t[:-1]
        if t.startswith("(") and t.endswith(")"):
            t = "-" + t[1:-1]
        try:
            return float(t)
        except ValueError:
            return None
    return None


DATE_FORMATS = ("%d/%m/%Y", "%m/%d/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d.%m.%Y", "%d %b %Y", "%d-%b-%Y", "%d/%m/%y")


def to_date(s):
    if isinstance(s, str):
        t = s.strip()
        for f in DATE_FORMATS:
            try:
                return datetime.strptime(t, f), f
            except ValueError:
                continue
    return None, None


def raw_kind(v) -> str:
    if isinstance(v, bool):
        return "boolean"
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return "number"
    if isinstance(v, (datetime, date, pd.Timestamp)):
        return "date"
    if isinstance(v, time):
        return "time"
    return "text"


# --------------------------------------------------------------------------- structure

def detect_header_row(rows: list[tuple]) -> int:
    """Return 0-based index of the most likely header row among the first rows."""
    best, best_score = 0, -1.0
    for i, row in enumerate(rows):
        vals = [v for v in row if v not in (None, "")]
        if len(vals) < 2:
            continue
        str_share = sum(isinstance(v, str) for v in vals) / len(vals)
        nxt = rows[i + 1] if i + 1 < len(rows) else ()
        has_data_below = sum(v not in (None, "") for v in nxt) >= max(2, len(vals) // 2)
        score = len(vals) * str_share + (2 if has_data_below else 0)
        if str_share >= 0.7 and score > best_score:
            best, best_score = i, score
    return best


def workbook_structure(path: Path) -> dict:
    """Workbook-level facts that pandas hides: hidden sheets, merged cells, formulas."""
    info = {}
    wb = load_workbook(path, data_only=False)
    for ws in wb.worksheets:
        formulas = 0
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    formulas += 1
        head = [tuple(r) for r in ws.iter_rows(min_row=1, max_row=HEADER_SCAN_ROWS, values_only=True)]
        info[ws.title] = {
            "hidden": ws.sheet_state != "visible",
            "dimensions": ws.dimensions,
            "merged_ranges": [str(m) for m in ws.merged_cells.ranges][:50],
            "formula_cells": formulas,
            "header_row_index": detect_header_row(head),
        }
    wb.close()
    return info


def csv_structure(path: Path) -> dict:
    head = pd.read_csv(path, header=None, nrows=HEADER_SCAN_ROWS, dtype=object)
    rows = [tuple(None if pd.isna(v) else v for v in r) for r in head.itertuples(index=False)]
    return {"(csv)": {"hidden": False, "dimensions": None, "merged_ranges": [], "formula_cells": 0,
                      "header_row_index": detect_header_row(rows)}}


def read_sheet(path: Path, sheet: str, header_idx: int) -> pd.DataFrame:
    if path.suffix.lower() == ".csv":
        df = pd.read_csv(path, header=header_idx, dtype=object)
    else:
        df = pd.read_excel(path, sheet_name=sheet, header=header_idx, dtype=object, engine="openpyxl")
    df = df.dropna(axis=1, how="all")
    df.columns = [str(c).strip() for c in df.columns]
    return df


# --------------------------------------------------------------------------- column profiling

def profile_column(name: str, s: pd.Series) -> dict:
    n = len(s)
    nn = s.dropna()
    nn = nn[~nn.map(lambda v: isinstance(v, str) and v.strip() == "")]
    kinds = Counter(raw_kind(v) for v in nn)
    col = {"name": name, "null_pct": round(1 - len(nn) / n, 4) if n else 1.0,
           "distinct": int(nn.astype(str).nunique()), "raw_types": dict(kinds)}
    if nn.empty:
        col.update(inferred_type="empty", suggested_role="unused")
        return col

    text_vals = [v for v in nn if isinstance(v, str)]
    # digit strings with leading zeros are codes (account numbers, cost centres), not numbers
    code_like = bool(text_vals) and kinds.get("number", 0) == 0 and \
        all(re.fullmatch(r"\d+", v.strip()) for v in text_vals) and \
        any(len(v.strip()) > 1 and v.strip().startswith("0") for v in text_vals)
    num_from_text = [] if code_like else [x for x in (to_number(v) for v in text_vals) if x is not None]
    parsed_dates = [to_date(v) for v in text_vals]
    date_from_text = [d for d, _ in parsed_dates if d is not None]
    date_fmts = Counter(f for _, f in parsed_dates if f)

    numeric_total = kinds.get("number", 0) + len(num_from_text)
    date_total = kinds.get("date", 0) + len(date_from_text)
    total = len(nn)

    if date_total / total >= 0.9:
        itype = "date"
    elif numeric_total / total >= 0.9:
        nums = [float(v) for v in nn if raw_kind(v) == "number"] + num_from_text
        itype = "integer" if all(float(x).is_integer() for x in nums) else "decimal"
    elif kinds.get("boolean", 0) / total >= 0.9:
        itype = "boolean"
    else:
        itype = "text"
    col["inferred_type"] = itype
    col["mixed_types"] = len(kinds) > 1
    if code_like:
        col["numeric_code"] = True   # keep as text: leading zeros matter
    if itype in ("integer", "decimal") and num_from_text:
        col["numeric_as_text"] = len(num_from_text)
    if itype == "date" and date_from_text:
        col["dates_as_text"] = len(date_from_text)
        col["text_date_formats"] = dict(date_fmts)
    col["sample_values"] = [jsonable(v) for v in list(dict.fromkeys(nn.tolist()))[:5]]

    if itype in ("integer", "decimal"):
        nums = pd.Series([float(v) for v in nn if raw_kind(v) == "number"] + num_from_text)
        q1, q3 = nums.quantile(0.25), nums.quantile(0.75)
        iqr = q3 - q1
        lo, hi = q1 - 3 * iqr, q3 + 3 * iqr
        outl = nums[(nums < lo) | (nums > hi)] if iqr > 0 else nums.iloc[0:0]
        col.update(min=jsonable(nums.min()), max=jsonable(nums.max()), mean=round(float(nums.mean()), 4),
                   sum=round(float(nums.sum()), 4), negatives=int((nums < 0).sum()),
                   outliers={"count": int(len(outl)), "examples": [jsonable(x) for x in outl.head(5)]})
    elif itype == "date":
        ds = pd.to_datetime(pd.Series([v for v in nn if raw_kind(v) == "date"] + date_from_text), errors="coerce")
        col.update(min=jsonable(ds.min()), max=jsonable(ds.max()))
    elif itype == "text" and col["distinct"] <= 500:
        groups: dict[str, set] = {}
        for v in text_vals:
            groups.setdefault(re.sub(r"\s+", " ", v).strip().casefold(), set()).add(v)
        variants = {k: sorted(vs) for k, vs in groups.items() if len(vs) > 1}
        if variants:
            col["inconsistent_categories"] = variants
        col["top_values"] = {str(k): int(v) for k, v in nn.astype(str).value_counts().head(10).items()}
        if col["distinct"] <= 50:
            col["values"] = sorted({str(v) for v in text_vals})
        ws_issues = sum(1 for v in text_vals if v != v.strip())
        if ws_issues:
            col["leading_trailing_spaces"] = ws_issues

    col["suggested_role"] = suggest_role(name, col, total)
    return col


def suggest_role(name: str, col: dict, total: int) -> str:
    t, d = col["inferred_type"], col["distinct"]
    uniq = d / total if total else 0
    if t == "date":
        return "date"
    if col.get("numeric_code") or (ID_NAME_PATTERN.search(name) and (uniq > 0.3 or t == "text")):
        return "identifier"
    if t in ("integer", "decimal"):
        return "measure"
    if t == "text":
        if uniq > 0.9 and total > 20:
            return "identifier"
        return "dimension" if d <= max(50, 0.2 * total) else "free_text"
    return "attribute"


# --------------------------------------------------------------------------- sheet profiling

def profile_sheet(path: Path, sheet: str, struct: dict) -> dict:
    hidx = struct["header_row_index"]
    df = read_sheet(path, sheet, hidx)
    excel_row = lambda i: hidx + 2 + i  # noqa: E731  (1-based Excel row number for df index i)

    blank = df.index[df.isna().all(axis=1)].tolist()
    total_rows, total_col = [], None
    for c in df.columns:
        m = df[c].map(lambda v: isinstance(v, str) and bool(TOTAL_PATTERN.search(v)))
        if m.any():
            total_rows += df.index[m].tolist()
            total_col = total_col or c
    total_rows = sorted(set(total_rows))
    clean = df.drop(index=set(blank) | set(total_rows))
    dup_mask = clean.astype(str).duplicated(keep="first")

    columns = [profile_column(c, clean[c]) for c in clean.columns]
    period_cols = [c for c in clean.columns if is_period_column(c)]
    wide = len(period_cols) >= 3
    wide_kind = ("tenor_buckets" if sum(bool(TENOR_PATTERN.match(c)) for c in period_cols) > len(period_cols) / 2
                 else "periods") if wide else None
    total_cols = [c for c in clean.columns if wide and TOTAL_PATTERN.search(str(c))]

    sheet_profile = {
        "name": sheet,
        "hidden": struct["hidden"],
        "dimensions": struct["dimensions"],
        "merged_ranges": struct["merged_ranges"],
        "formula_cells": struct["formula_cells"],
        "header_row": hidx + 1,
        "rows_skipped_above_header": hidx,
        "n_rows_raw": int(len(df)),
        "n_rows_clean": int(len(clean)),
        "n_cols": int(len(df.columns)),
        "blank_rows": [excel_row(i) for i in blank][:100],
        "total_rows": [excel_row(i) for i in total_rows][:100],
        "total_rows_column": total_col,
        "duplicate_rows": int(dup_mask.sum()),
        "wide_format": {"detected": wide, "kind": wide_kind, "period_columns": period_cols,
                        "total_columns": total_cols,
                        "id_columns": [c for c in clean.columns if c not in period_cols and c not in total_cols]
                        if wide else []},
        "columns": columns,
    }
    sheet_profile["issues"] = list_issues(sheet_profile)
    sheet_profile["suggested_cleaning_rules"] = suggest_rules(sheet_profile)
    return sheet_profile


def list_issues(sp: dict) -> list[str]:
    out = []
    if sp["hidden"]:
        out.append("Sheet is hidden - confirm whether it should be used.")
    if sp["rows_skipped_above_header"]:
        out.append(f"Header is on row {sp['header_row']}; {sp['rows_skipped_above_header']} title/notes row(s) above it.")
    if sp["total_rows"]:
        out.append(f"{len(sp['total_rows'])} subtotal/total row(s) mixed into the data (column '{sp['total_rows_column']}').")
    if sp["blank_rows"]:
        out.append(f"{len(sp['blank_rows'])} blank separator row(s).")
    if sp["duplicate_rows"]:
        out.append(f"{sp['duplicate_rows']} exact duplicate row(s).")
    if sp["formula_cells"]:
        out.append(f"{sp['formula_cells']} formula cell(s); values may be missing if the file was never recalculated in Excel.")
    if sp["wide_format"]["detected"]:
        kind = "tenor bucket" if sp["wide_format"].get("kind") == "tenor_buckets" else "period"
        out.append(f"Wide format: {len(sp['wide_format']['period_columns'])} {kind} columns should be unpivoted.")
    for c in sp["columns"]:
        if c.get("numeric_as_text"):
            out.append(f"'{c['name']}': {c['numeric_as_text']} number(s) stored as text.")
        if c.get("dates_as_text"):
            out.append(f"'{c['name']}': {c['dates_as_text']} date(s) stored as text {list(c.get('text_date_formats', {}))}.")
        if c.get("inconsistent_categories"):
            out.append(f"'{c['name']}': inconsistent spellings for {len(c['inconsistent_categories'])} value(s).")
        if c.get("outliers", {}).get("count"):
            out.append(f"'{c['name']}': {c['outliers']['count']} extreme outlier(s), e.g. {c['outliers']['examples'][:3]}.")
        if c["null_pct"] > 0.2 and c["inferred_type"] != "empty":
            out.append(f"'{c['name']}': {fmt_pct(c['null_pct'])} missing.")
    return out


def suggest_rules(sp: dict) -> list[dict]:
    """Draft cleaning rules in the format tools/load_data.py executes. The agent reviews them."""
    rules = []
    if sp["blank_rows"]:
        rules.append({"op": "drop_blank_rows"})
    if sp["total_rows"]:
        rules.append({"op": "drop_rows_matching", "column": sp["total_rows_column"],
                      "pattern": TOTAL_PATTERN.pattern, "reason": "subtotal/total rows"})
    for c in sp["columns"]:
        if c.get("leading_trailing_spaces") and not c.get("inconsistent_categories"):
            rules.append({"op": "strip_whitespace", "columns": [c["name"]]})
        if c.get("inconsistent_categories") and c.get("values"):
            top = c.get("top_values", {})
            groups: dict[str, list[str]] = {}
            for v in c["values"]:
                groups.setdefault(re.sub(r"\s+", " ", v).strip().casefold(), []).append(v)
            canonical = sorted(max(vs, key=lambda v: top.get(v, 0)).strip() for vs in groups.values())
            rules.append({"op": "canonical_values", "column": c["name"], "values": canonical,
                          "reason": "standardise spellings; future case/space variants are handled too",
                          "needs_confirmation": True})
        if c.get("numeric_as_text") or (c["inferred_type"] in ("integer", "decimal") and c.get("mixed_types")):
            rules.append({"op": "to_numeric", "columns": [c["name"]]})
        if c["inferred_type"] == "date":
            fmts = list(c.get("text_date_formats", {}))
            rules.append({"op": "to_datetime", "columns": [c["name"]],
                          "dayfirst": not fmts or fmts[0].startswith("%d")})
    if sp["duplicate_rows"]:
        rules.append({"op": "drop_duplicates", "needs_confirmation": True,
                      "reason": "exact duplicates may be genuine repeat transactions"})
    if sp["wide_format"]["detected"]:
        if sp["wide_format"].get("total_columns"):
            rules.append({"op": "drop_columns", "columns": sp["wide_format"]["total_columns"],
                          "reason": "total column in a wide sheet - recompute totals from the buckets"})
        rules.append({"op": "unpivot", "id_columns": sp["wide_format"]["id_columns"],
                      "value_columns": "periods",  # every period column, including future months
                      "var_name": "Tenor bucket" if sp["wide_format"].get("kind") == "tenor_buckets" else "Period",
                      "value_name": "Value", "needs_confirmation": True})
    return rules


# --------------------------------------------------------------------------- drift / caching

def _stable_columns(sheet: dict) -> list[str]:
    """Column names excluding period columns of wide sheets (new months are expected growth)."""
    periods = set(sheet.get("wide_format", {}).get("period_columns", []))
    return [c["name"] for c in sheet["columns"] if c["name"] not in periods]


def structure_signature(sheets: list[dict]) -> str:
    sig = [(s["name"], s["header_row"], _stable_columns(s), s.get("wide_format", {}).get("detected", False))
           for s in sheets]
    return hashlib.sha256(json.dumps(sig).encode()).hexdigest()[:16]


def compare(old: dict, new: dict) -> dict:
    o = {s["name"]: s for s in old.get("sheets", [])}
    n = {s["name"]: s for s in new["sheets"]}
    # recompute both signatures so profiles made by older versions of this script compare fairly
    drift = {"structure_changed": structure_signature(old["sheets"]) != structure_signature(new["sheets"]),
             "sheets_added": sorted(set(n) - set(o)), "sheets_removed": sorted(set(o) - set(n)), "sheets": {}}
    for name in set(o) & set(n):
        oc = {c["name"]: c for c in o[name]["columns"] if c["name"] in _stable_columns(o[name])}
        nc = {c["name"]: c for c in n[name]["columns"] if c["name"] in _stable_columns(n[name])}
        d = {"columns_added": sorted(set(nc) - set(oc)), "columns_removed": sorted(set(oc) - set(nc)),
             "type_changes": {k: [oc[k]["inferred_type"], nc[k]["inferred_type"]]
                              for k in set(oc) & set(nc) if oc[k]["inferred_type"] != nc[k]["inferred_type"]},
             "header_row_moved": [o[name]["header_row"], n[name]["header_row"]]
             if o[name]["header_row"] != n[name]["header_row"] else None,
             "row_count": [o[name]["n_rows_clean"], n[name]["n_rows_clean"]]}
        op_ = set(o[name].get("wide_format", {}).get("period_columns", []))
        np_ = set(n[name].get("wide_format", {}).get("period_columns", []))
        if op_ != np_:
            drift.setdefault("period_columns", {})[name] = {"added": sorted(np_ - op_), "removed": sorted(op_ - np_)}
        if any([d["columns_added"], d["columns_removed"], d["type_changes"], d["header_row_moved"]]):
            drift["sheets"][name] = d
    return drift


def empty_annotations(sheets: list[dict]) -> dict:
    """Filled in by the Profiler agent and confirmed by a human."""
    return {
        "status": "draft",            # draft -> confirmed
        "confirmed_by": None,
        "confirmed_on": None,
        "sheets": {s["name"]: {
            "include": not s["hidden"],
            "description": "",        # what this sheet is, in business terms
            "grain": "",              # what ONE row represents
            "column_roles": {},       # column -> dimension|measure|date|identifier|attribute|unused
            "column_meanings": {},    # column -> business meaning, units, currency, VAT etc.
            "column_additivity": {},  # measure -> additive | semi_additive over <dim> | non_additive
            "column_units": {},       # measure -> unit, e.g. "EUR", "%", "bp", "days" (goes into the catalog)
            "prefilled_from_catalog": {},  # column -> catalog id (filled automatically from knowledge/)
            "cleaning_rules": [],     # confirmed rules (copied/edited from suggested_cleaning_rules)
        } for s in sheets},
        "domains": [],                # confirmed domain packs, e.g. ["treasury_common", "money_market"]
        "feasible_analyses": [],
        "not_feasible": [],           # e.g. "Margin - no cost data"
        "open_questions": [],
        # answered questions move here instead of being deleted, so the Curator can turn them into
        # decisions without the chat: {"question", "answer", "by", "on": "YYYY-MM-DD"}
        "answers": [],
    }


# --------------------------------------------------------------------------- draft annotations

NON_ADDITIVE_NAME = re.compile(r"(?i)rate|%|ratio|yield|price|spread|margin|spot|percent|\bbps?\b|coupon|utili[sz]")
STOCK_NAME = re.compile(r"(?i)balance|outstanding|notional|principal|nominal|position|exposure|limit|"
                        r"holding|inventory|market value|book value|mtm")
FLOW_NAME = re.compile(r"(?i)amount|payment|interest|flow|volume|revenue|cost|fee|qty|quantity|sales|paid|received")
DATE_IN_NAME = re.compile(r"(?<!\d)(19|20)\d{2}(?:[-_.]?(0[1-9]|1[0-2])(?:[-_.]?(0[1-9]|[12]\d|3[01]))?)?(?!\d)")


def recurring_hint(profile: dict) -> str | None:
    """A dated file name with no logical source is probably one of a series: suggest a pattern."""
    name = profile["source"]["file"]
    if profile["source"].get("logical_sources") or not DATE_IN_NAME.search(name):
        return None
    return DATE_IN_NAME.sub("*", name, count=1)


def _guess_additivity(name: str) -> str | None:
    for rx, value in ((NON_ADDITIVE_NAME, "non_additive"), (STOCK_NAME, "semi_additive over time"),
                      (FLOW_NAME, "additive")):
        if rx.search(name):
            return value
    return None


def prefill_annotations(profile: dict) -> list[str]:
    """Draft the mechanical parts of the annotations from what the profiler measured, so the
    agent writes only business meaning. Fills gaps only (catalog and human entries win), turns
    every guess into an open question, and never changes the status: a person still confirms."""
    ann = profile["annotations"]
    if ann.get("status") == "confirmed":
        return ["annotations already confirmed - not prefilled"]
    questions = ann.setdefault("open_questions", [])
    done = []

    def ask(q: str):
        if q not in questions:
            questions.append(q)

    for s in profile["sheets"]:
        sa = ann["sheets"].setdefault(s["name"], {})
        if sa.get("include") is False:
            continue
        roles = sa.setdefault("column_roles", {})
        meanings = sa.setdefault("column_meanings", {})
        additivity = sa.setdefault("column_additivity", {})
        n_roles = n_add = 0
        wide = s.get("wide_format", {})
        reshaped = set(wide.get("period_columns", [])) | set(wide.get("total_columns", []))
        if wide.get("period_columns"):
            cols = wide["period_columns"]
            ask(f"[{s['name']}] {len(cols)} period columns ({cols[0]} .. {cols[-1]}) will be unpivoted into one "
                "value column - what does the value mean, and is it a flow (additive) or a stock?")
        for c in s["columns"]:
            col = c["name"]
            if col in reshaped:
                continue
            if col not in roles:
                roles[col] = c["suggested_role"]
                n_roles += 1
            if roles[col] == "measure" and col not in additivity:
                guess = _guess_additivity(col)
                if guess:
                    additivity[col] = guess
                    n_add += 1
                    ask(f"[{s['name']}] '{col}': additivity guessed as '{guess}' from its name - confirm.")
                else:
                    ask(f"[{s['name']}] '{col}' looks numeric - is it a measure (additive / stock / rate) "
                        "or an identifier/code?")
            if c.get("distinct") == 1 and roles[col] in ("date", "dimension") and c.get("null_pct", 0) == 0:
                value = str((c.get("sample_values") or ["?"])[0])[:10] if roles[col] == "date" else \
                    str((c.get("sample_values") or ["?"])[0])
                meanings.setdefault(col, f"Single value in this file: {value}.")
                ask(f"[{s['name']}] '{col}' has one value ({value}) - "
                    + ("is it the as-of date for every figure?" if roles[col] == "date"
                       else "expected (e.g. a single-currency or single-entity file)?"))
            if roles[col] == "measure" and c.get("outliers", {}).get("count"):
                ask(f"[{s['name']}] '{col}': {c['outliers']['count']} statistical outlier(s), e.g. "
                    f"{c['outliers']['examples'][:3]} - real values to keep, or errors?")
        if not sa.get("cleaning_rules") and s.get("suggested_cleaning_rules"):
            sa["cleaning_rules"] = copy.deepcopy(s["suggested_cleaning_rules"])
            done.append(f"[{s['name']}] {len(sa['cleaning_rules'])} cleaning rule(s) copied from suggestions")
        if not sa.get("grain"):
            ids = [c["name"] for c in s["columns"] if roles.get(c["name"]) == "identifier"
                   and c.get("distinct") == s["n_rows_clean"]]
            if ids:
                sa["grain"] = f"one row per {ids[0]} (draft - confirm)"
        done.append(f"[{s['name']}] {n_roles} role(s), {n_add} additivity guess(es) drafted")
    pattern = recurring_hint(profile)
    if pattern:
        ask(f"Is this file one of a recurring series? If yes, register a logical source (pattern "
            f"'{pattern}') BEFORE scaffolding a report, so reports read the current file.")
    ann["prefilled_by"] = f"profile_excel.py --prefill on {date.today().isoformat()}"
    return done


# --------------------------------------------------------------------------- HTML profile report

def write_profile_html(profile: dict, out: Path):
    src = profile["source"]
    r = Report(f"Data profile: {src['file']}", subtitle="What the profiler found in this workbook "
               "before any analysis is done.", sources=[src["file"]])
    sheets = profile["sheets"]
    n_issues = sum(len(s["issues"]) for s in sheets)
    kn = profile.get("knowledge", {})
    n_cols = kn.get("recognised", 0) + kn.get("unknown", 0) + kn.get("ambiguous", 0)
    r.headline(f"{len(sheets)} sheet(s), {sum(s['n_rows_clean'] for s in sheets):,} usable rows, "
               f"{n_issues} issue(s) to resolve before analysis.",
               (f"{kn.get('recognised', 0)} of {n_cols} columns recognised from the knowledge base. " if kn else "")
               + ("Domain packs: " + ", ".join(profile["annotations"].get("domains") or []) + ". "
                  if profile["annotations"].get("domains") else
                  ("Suggested domain packs: " + ", ".join(d["id"] for d in profile.get("suggested_domains", [])) + ". "
                   if profile.get("suggested_domains") else ""))
               + "Status: " + profile["annotations"]["status"] +
               (" - " + ", ".join(profile["drift"]["sheets"]) + " changed since last profile"
                if profile.get("drift", {}).get("sheets") else ""))
    for s in sheets:
        ann = profile["annotations"]["sheets"].get(s["name"], {})
        r.section(f"Sheet: {s['name']}" + (" (hidden)" if s["hidden"] else ""),
                  ann.get("description") or "Business description pending from the Profiler agent.")
        r.kpis([{"label": "Usable rows", "value": fmt_num(s["n_rows_clean"])},
                {"label": "Columns", "value": fmt_num(s["n_cols"])},
                {"label": "Header row", "value": str(s["header_row"])},
                {"label": "Issues", "value": str(len(s["issues"]))}])
        if ann.get("grain"):
            r.callout(ann["grain"], title="One row represents")
        if s["issues"]:
            r.findings(s["issues"])
        def known(c):
            k = c.get("catalog", {})
            if k.get("match") == "confirmed":
                return k.get("label") or k["id"]
            if k.get("match") == "ambiguous":
                return "Ambiguous: " + " / ".join(k["candidates"])
            return "New - needs definition"
        r.table([{"Column": c["name"], "Known as": known(c), "Type": c["inferred_type"],
                  "Role": ann.get("column_roles", {}).get(c["name"], c["suggested_role"]),
                  "Missing": fmt_pct(c["null_pct"]), "Distinct": c["distinct"],
                  "Examples": ", ".join(str(v).replace("T00:00:00", "") for v in c.get("sample_values", [])[:3])[:60]}
                 for c in s["columns"]], title="Columns")
    qs = profile["annotations"].get("open_questions", [])
    if qs:
        r.section("Open questions for the business")
        r.findings(qs)
    answered = profile["annotations"].get("answers", [])
    if answered:
        r.section("Questions answered", "What the business confirmed about this file.")
        r.findings(f"{a.get('question', '')} - {a.get('answer', '')} ({a.get('by') or 'unknown'}, {a.get('on') or 'undated'})"
                   for a in answered)
    r.methodology([f"Profiled {src['profiled_at'][:16].replace('T', ' ')} with tools/profile_excel.py",
                   f"File hash {src['sha256'][:12]}, structure signature {src['structure_signature']}"])
    r.save(out)


# --------------------------------------------------------------------------- relationships

def find_relationships(profiles: list[dict]) -> list[dict]:
    cands = []
    for p in profiles:
        path = ROOT / p["source"]["path"]
        for s in p["sheets"]:
            for c in s["columns"]:
                if c["suggested_role"] in ("identifier", "dimension") and c["distinct"] > 1:
                    cands.append((p["source"]["file"], s["name"], c["name"], path, s["header_row"] - 1))
    values = {}
    for f, sh, col, path, h in cands:
        try:
            df = read_sheet(path, sh, h)
            values[(f, sh, col)] = set(df[col].dropna().astype(str).str.strip().str.casefold())
        except Exception:  # noqa: BLE001
            continue
    norm = lambda x: re.sub(r"[^a-z0-9]", "", x.lower())  # noqa: E731
    rels = []
    keys = list(values)
    for i, a in enumerate(keys):
        for b in keys[i + 1:]:
            if a[:2] == b[:2]:
                continue
            same_name = norm(a[2]) == norm(b[2])
            va, vb = values[a], values[b]
            if not va or not vb:
                continue
            overlap = len(va & vb)
            rate_a, rate_b = overlap / len(va), overlap / len(vb)
            if (same_name and overlap) or max(rate_a, rate_b) >= 0.8 and min(len(va), len(vb)) >= 5:
                rels.append({"left": {"file": a[0], "sheet": a[1], "column": a[2]},
                             "right": {"file": b[0], "sheet": b[1], "column": b[2]},
                             "same_name": same_name, "left_match_rate": round(rate_a, 3),
                             "right_match_rate": round(rate_b, 3)})
    return sorted(rels, key=lambda r: -max(r["left_match_rate"], r["right_match_rate"]))


# --------------------------------------------------------------------------- main

def rule_gaps(profile: dict) -> list[str]:
    """Values in the data that confirmed value rules do not cover (e.g. a new region or spelling).
    Used when annotations are carried forward, so a refresh stops at profiling, not at checks."""
    gaps = []
    key = lambda v: re.sub(r"\s+", " ", str(v)).strip().casefold()  # noqa: E731
    for s in profile["sheets"]:
        ann = profile["annotations"]["sheets"].get(s["name"], {})
        if ann.get("include") is False:
            continue
        cols = {c["name"]: c for c in s["columns"]}
        for r in ann.get("cleaning_rules", []):
            c = cols.get(r.get("column"))
            if not c or "values" not in c:
                continue
            if r["op"] == "canonical_values":
                allowed = {key(v) for v in r["values"]}
                new = sorted({v for v in c["values"] if key(v) not in allowed})
                if new:
                    gaps.append(f"{s['name']}.{c['name']}: values not in the confirmed list: {new}")
            elif r["op"] == "map_values":
                known = {key(v) for v in list(r["map"]) + list(r["map"].values())}
                targets = set(r["map"].values())
                new = sorted({v for v in c["values"] if v not in r["map"] and v not in targets
                              and key(v) in known})
                if new:
                    gaps.append(f"{s['name']}.{c['name']}: new spelling variants not in map_values: {new}")
    return gaps


def _apply_domains(profile: dict):
    """Rank domain packs for this file; take confirmed packs from its logical source if not set."""
    try:
        profile["suggested_domains"] = suggest_domains(profile)
    except Exception as e:  # noqa: BLE001 - a malformed pack must never stop profiling
        profile["suggested_domains"] = []
        print(f"  (domain suggestion skipped: {e})")
    ann = profile["annotations"]
    if not ann.get("domains"):
        ann["domains"] = recorded_domains(profile)


def predecessor_profile(path: Path) -> dict | None:
    """Most recent existing profile of an earlier file from the same logical source."""
    for sid in match_source(path.name):
        earlier = [f for f in source_files(sid) if f.name != path.name]
        earlier = [f for f in earlier if f.name.lower() < path.name.lower()] or earlier
        for f in reversed(earlier):
            pj = PROFILE_DIR / f"{f.stem}.profile.json"
            if pj.exists():
                return json.loads(pj.read_text(encoding="utf-8"))
    return None


def profile_file(path: Path, force: bool) -> tuple[dict, str]:
    PROFILE_DIR.mkdir(exist_ok=True)
    out_json = PROFILE_DIR / f"{path.stem}.profile.json"
    digest = sha256(path)
    old = json.loads(out_json.read_text()) if out_json.exists() else None
    if old and not force and old["source"]["sha256"] == digest:
        old["source"]["logical_sources"] = match_source(path.name)
        old["source"]["structure_signature"] = structure_signature(old["sheets"])
        k = apply_catalog(old)  # the knowledge base may have grown since last run
        _apply_domains(old)
        out_json.write_text(json.dumps(old, indent=2, default=jsonable), encoding="utf-8")
        write_profile_html(old, PROFILE_DIR / f"{path.stem}.profile.html")
        return old, f"unchanged (cached profile reused); knowledge: {k['recognised']} recognised, " \
                    f"{k['unknown']} unknown, {k['ambiguous']} ambiguous"

    struct = csv_structure(path) if path.suffix.lower() == ".csv" else workbook_structure(path)
    sheets = [profile_sheet(path, name, st) for name, st in struct.items()]
    profile = {
        "schema_version": SCHEMA_VERSION,
        "source": {"file": path.name, "path": str(path.relative_to(ROOT)).replace("\\", "/"),
                   "sha256": digest, "size_bytes": path.stat().st_size,
                   "profiled_at": datetime.now().isoformat(timespec="seconds"),
                   "structure_signature": structure_signature(sheets),
                   "logical_sources": match_source(path.name)},
        "sheets": sheets,
    }
    if old:
        profile["drift"] = compare(old, profile)
        profile["annotations"] = old.get("annotations") or empty_annotations(sheets)
        profile["annotations"].pop("review_reasons", None)
        if profile["drift"]["structure_changed"]:
            profile["annotations"]["status"] = "needs_review"
            status = "STRUCTURE CHANGED - annotations need review"
        else:
            status = "data changed, structure unchanged - annotations carried forward"
            gaps = rule_gaps(profile)
            if gaps and profile["annotations"].get("status") == "confirmed":
                profile["annotations"]["status"] = "needs_review"
                profile["annotations"]["review_reasons"] = gaps
                status = "DATA NEEDS REVIEW - " + "; ".join(gaps)
    elif (pred := predecessor_profile(path)) is not None:
        # a new file of a known recurring source: inherit the confirmed annotations
        profile["drift"] = compare(pred, profile)
        profile["annotations"] = copy.deepcopy(pred["annotations"])
        profile["annotations"]["inherited_from"] = pred["source"]["file"]
        profile["annotations"].pop("review_reasons", None)
        profile["annotations"].pop("prefilled_from_catalog", None)
        if profile["drift"]["structure_changed"] or pred["annotations"].get("status") != "confirmed":
            profile["annotations"]["status"] = "needs_review"
            status = (f"NEW FILE of source {profile['source']['logical_sources']} - structure differs from "
                      f"{pred['source']['file']}, annotations need review")
        else:
            status = (f"new file of source {profile['source']['logical_sources']}; same structure as "
                      f"{pred['source']['file']} - confirmed annotations inherited")
            gaps = rule_gaps(profile)
            if gaps:
                profile["annotations"]["status"] = "needs_review"
                profile["annotations"]["review_reasons"] = gaps
                status += " BUT DATA NEEDS REVIEW - " + "; ".join(gaps)
    else:
        profile["drift"] = {}
        profile["annotations"] = empty_annotations(sheets)
        status = "new profile"
    k = apply_catalog(profile)
    _apply_domains(profile)
    status += f"; knowledge: {k['recognised']} recognised, {k['unknown']} unknown, {k['ambiguous']} ambiguous"
    out_json.write_text(json.dumps(profile, indent=2, default=jsonable), encoding="utf-8")
    write_profile_html(profile, PROFILE_DIR / f"{path.stem}.profile.html")
    return profile, status


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("files", nargs="*", help="files to profile (default: all in data/raw)")
    ap.add_argument("--force", action="store_true", help="re-profile even if unchanged")
    ap.add_argument("--html-only", action="store_true", help="re-render HTML from existing profiles")
    ap.add_argument("--prefill", action="store_true",
                    help="draft roles, additivity, cleaning rules, grain and open questions in unconfirmed profiles")
    ap.add_argument("--quiet", action="store_true", help="one status line per file (issues are in the profile)")
    args = ap.parse_args()
    if args.html_only:
        for pj in sorted(PROFILE_DIR.glob("*.profile.json")):
            write_profile_html(json.loads(pj.read_text(encoding="utf-8")), pj.with_suffix("").with_suffix(".profile.html"))
            print(f"Re-rendered {pj.stem}.html")
        return
    files = [Path(f).resolve() for f in args.files] or sorted(
        p for p in RAW_DIR.iterdir() if p.suffix.lower() in SUPPORTED and not p.name.startswith("~$"))
    if not files:
        print(f"No supported files found in {RAW_DIR}")
        return
    profiles = []
    for f in files:
        if f.suffix.lower() not in SUPPORTED:
            print(f"SKIP {f.name}: unsupported type (convert .xls to .xlsx)")
            continue
        p, status = profile_file(f, args.force)
        profiles.append(p)
        if args.quiet and not args.prefill:
            print(f"{f.name}: {status}")
            continue
        print(f"\n== {f.name}: {status}")
        if args.prefill:
            for line in prefill_annotations(p):
                print(f"  prefill: {line}")
            if p["annotations"].get("status") != "confirmed":
                (PROFILE_DIR / f"{f.stem}.profile.json").write_text(json.dumps(p, indent=2, default=jsonable),
                                                                    encoding="utf-8")
                write_profile_html(p, PROFILE_DIR / f"{f.stem}.profile.html")
                print(f"  prefill: {len(p['annotations']['open_questions'])} open question(s) to ask the user")
        pattern = recurring_hint(p)
        if pattern and p["annotations"].get("status") != "confirmed":
            print(f"  recurring? dated name, no logical source - if it is a series, register pattern '{pattern}'")
        for s in p["sheets"]:
            print(f"  [{s['name']}] header row {s['header_row']}, {s['n_rows_clean']} usable rows, "
                  f"{s['n_cols']} cols, {len(s['issues'])} issue(s)")
            if s.get("unknown_columns"):
                print(f"    not in knowledge base: {s['unknown_columns']}")
            for i in s["issues"]:
                print(f"    - {i}")
        doms = p["annotations"].get("domains") or []
        sugg = [f"{d['id']} ({d['score']})" for d in p.get("suggested_domains", []) if not d.get("as_parent_of")]
        print(f"  domains: {doms if doms else 'not set'}" + (f" | suggested: {', '.join(sugg)}" if sugg and not doms else ""))
    all_profiles = [json.loads(p.read_text()) for p in PROFILE_DIR.glob("*.profile.json")]
    # for recurring sources only the current file takes part (older months are history)
    current = {f.name for sid in {s for p in all_profiles for s in p["source"].get("logical_sources", [])}
               for f in source_files(sid)[-1:]}
    all_profiles = [p for p in all_profiles
                    if not p["source"].get("logical_sources") or p["source"]["file"] in current]
    rels = find_relationships(all_profiles)
    (PROFILE_DIR / "_relationships.json").write_text(json.dumps(rels, indent=2), encoding="utf-8")
    if args.quiet:
        return
    print(f"\nCandidate relationships: {len(rels)}")
    for r in rels[:5]:
        print(f"  {r['left']['file']}:{r['left']['sheet']}.{r['left']['column']} <-> "
              f"{r['right']['file']}:{r['right']['sheet']}.{r['right']['column']} "
              f"(match {r['left_match_rate']:.0%} / {r['right_match_rate']:.0%})")
    print(f"\nProfiles written to {PROFILE_DIR}")


if __name__ == "__main__":
    main()
