"""
Generate deliberately messy Treasury sample workbooks for smoke-testing a clean workspace.
Replace them with real departmental files for actual tests.

Creates in data/raw/:
  bank_balances_2026_09.xlsx  Liquidity: daily closing/available balances by entity, bank,
                              account and currency. Title rows above the header, a daily "Total"
                              row mixing currencies, account numbers with leading zeros,
                              currency case variants, some amounts stored as text, overdrafts.
  mm_deals_2026_09.xlsx       Money market: deal blotter with trade/value/maturity dates,
                              principal, rate in percent (some as text), day count, status
                              (incl. cancelled deals) and counterparty spelling variants.
  irr_gap_2026_q3.xlsx        Interest rate: repricing gap in EUR/USD millions, tenor buckets as
                              columns (O/N ... >5Y), subtotal rows and a Total column.

Usage:  python samples/make_treasury_samples.py
        python samples/make_treasury_samples.py --next-month   # bank_balances_2026_10.xlsx
"""
import random
import sys
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook

OUT = Path(__file__).resolve().parents[1] / "data" / "raw"
OUT.mkdir(parents=True, exist_ok=True)

ACCOUNTS = [  # entity, bank, account no, currency, starting balance
    ("TreasCo UK Ltd", "Bank A", "00012345", "GBP", 18_000_000),
    ("TreasCo UK Ltd", "Bank B", "00098761", "GBP", 6_500_000),
    ("TreasCo UK Ltd", "Bank A", "00012399", "EUR", 4_200_000),
    ("TreasCo GmbH", "Bank C", "04400123", "EUR", 22_000_000),
    ("TreasCo GmbH", "Bank A", "04400877", "USD", 3_100_000),
    ("TreasCo Inc", "Bank D", "70000451", "USD", 31_000_000),
    ("TreasCo Inc", "Bank B", "70000999", "USD", 900_000),
    ("TreasCo NV", "Bank C", "05500321", "EUR", 1_200_000),
]


def business_days(year, month):
    d = date(year, month, 1)
    while d.month == month:
        if d.weekday() < 5:
            yield d
        d += timedelta(days=1)


def make_balances(year=2026, month=9, seed=11):
    random.seed(seed)
    wb = Workbook()
    ws = wb.active
    ws.title = "Balances"
    ws["A1"] = f"Daily bank balances - {date(year, month, 1):%B %Y}"
    ws.merge_cells("A1:G1")
    ws["A2"] = "Amounts in account currency. Source: bank statements (MT940)."
    ws.append([])
    ws.append(["Value Date", "Entity", "Bank", "Account No", "Currency", "Closing Balance", "Available Balance"])
    bal = {a[2]: a[4] for a in ACCOUNTS}
    for d in business_days(year, month):
        day_total = 0.0
        for ent, bank, acc, ccy, _ in ACCOUNTS:
            bal[acc] = round(bal[acc] * (1 + random.uniform(-0.04, 0.04)) + random.uniform(-4e5, 4e5), 2)
            if acc == "70000999" and random.random() < 0.3:
                bal[acc] = round(-random.uniform(5e4, 3e5), 2)          # overdraft
            closing = bal[acc]
            available = round(closing + (500_000 if bank == "Bank B" else 0) - random.uniform(0, 2e4), 2)
            ccy_cell = ccy.lower() if random.random() < 0.05 else ccy
            closing_cell = f"{closing:,.2f}" if random.random() < 0.04 else closing
            ws.append([d, ent, bank, acc, ccy_cell, closing_cell, available])
            day_total += closing
        ws.append([d, "Total", None, None, None, round(day_total, 2), None])   # mixes currencies!
    name = f"bank_balances_{year}_{month:02d}.xlsx"
    wb.save(OUT / name)
    return name


def make_mm_deals(seed=5):
    random.seed(seed)
    wb = Workbook()
    ws = wb.active
    ws.title = "Deals"
    ws.append(["Deal ID", "Trade Date", "Value Date", "Maturity Date", "Counterparty", "Deal Type",
               "Currency", "Principal", "Rate (%)", "Day Count", "Interest", "Status"])
    cps = ["Bank A plc", "Bank C AG", "Bank D NA", "MMF Liquidity Fund", "Bank B Ltd"]
    basis = {"EUR": ("ACT/360", 360), "USD": ("ACT/360", 360), "GBP": ("ACT/365", 365)}
    base_rate = {"EUR": 2.05, "USD": 4.10, "GBP": 3.95}
    for i in range(64):
        ccy = random.choice(["EUR", "EUR", "USD", "GBP"])
        trade = date(2026, 8, 18) + timedelta(days=random.randint(0, 40))
        value = trade + timedelta(days=random.choice([0, 0, 1, 2]))
        tenor = random.choice([1, 7, 14, 30, 30, 61, 91, 91, 182])
        maturity = value + timedelta(days=tenor)
        dtype = random.choice(["Deposit", "Deposit", "Placement", "Borrowing"])
        principal = random.choice([5, 10, 10, 15, 20, 25, 50]) * 1_000_000
        rate = round(base_rate[ccy] + random.uniform(-0.25, 0.30) + (0.15 if dtype == "Borrowing" else 0), 3)
        dc, days_basis = basis[ccy]
        interest = round(principal * rate / 100 * tenor / days_basis, 2)
        cp = random.choice(cps)
        if cp == "Bank A plc" and random.random() < 0.3:
            cp = "BANK A PLC"
        status = "Cancelled" if random.random() < 0.05 else ("Matured" if maturity < date(2026, 9, 30) else "Live")
        rate_cell = f"{rate}%" if random.random() < 0.06 else rate
        trade_cell = trade.strftime("%d/%m/%Y") if random.random() < 0.04 else trade
        ws.append([f"MM{26000 + i}", trade_cell, value, maturity, cp, dtype, ccy, principal, rate_cell,
                   dc, interest, status])
    wb.save(OUT / "mm_deals_2026_09.xlsx")
    return "mm_deals_2026_09.xlsx"


def make_gap(seed=3):
    random.seed(seed)
    wb = Workbook()
    ws = wb.active
    ws.title = "Repricing Gap"
    ws["A1"] = "Interest rate repricing gap - 30 Sep 2026"
    ws["A2"] = "Amounts in millions. Repricing (not contractual) buckets. Derivatives excluded."
    ws.append([])
    buckets = ["O/N", "0-1M", "1-3M", "3-6M", "6-12M", "1-2Y", "2-5Y", ">5Y"]
    ws.append(["Line", "Currency", *buckets, "Total"])
    lines = {"Assets": ["Loans - floating", "Loans - fixed", "Securities", "Interbank placements"],
             "Liabilities": ["Customer deposits - sight", "Customer deposits - term", "Wholesale funding"]}
    for ccy, scale in (("EUR", 1.0), ("USD", 0.6)):
        totals = {}
        for side, items in lines.items():
            side_tot = [0.0] * len(buckets)
            for item in items:
                vals = [round(random.uniform(5, 120) * scale, 1) for _ in buckets]
                side_tot = [a + b for a, b in zip(side_tot, vals)]
                ws.append([item, ccy, *vals, round(sum(vals), 1)])
            side_tot = [round(v, 1) for v in side_tot]
            totals[side] = side_tot
            ws.append([f"Total {side.lower()}", ccy, *side_tot, round(sum(side_tot), 1)])
        gap = [round(a - l, 1) for a, l in zip(totals["Assets"], totals["Liabilities"])]
        ws.append(["Gap", ccy, *gap, round(sum(gap), 1)])
    wb.save(OUT / "irr_gap_2026_q3.xlsx")
    return "irr_gap_2026_q3.xlsx"


if __name__ == "__main__":
    if "--next-month" in sys.argv:
        print("Wrote", make_balances(2026, 10, seed=12))
        sys.exit(0)
    for f in (make_balances(), make_mm_deals(), make_gap()):
        print("Wrote", OUT / f)
